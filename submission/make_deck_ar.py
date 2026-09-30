# -*- coding: utf-8 -*-
"""Build the mizan Arabic 7-slide deck: PPTX + PDF from one content list.

RTL policy: PPTX keeps logical Arabic and sets pPr@rtl=1 (PowerPoint shapes it).
PDF reshapes with arabic_reshaper + python-bidi because reportlab cannot shape.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from reportlab.lib.pagesizes import landscape, A4
from reportlab.lib.colors import HexColor
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from PIL import Image, ImageDraw, ImageFont
import arabic_reshaper
from bidi.algorithm import get_display

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
FONTS = r"C:\Windows\Fonts"
TERMINAL_SRC = r"C:\Users\Saudi\AppData\Local\Temp\opencode\demo-out.txt"
TERMINAL_PNG = os.path.join(HERE, "demo-terminal.png")

GREEN = RGBColor(0x0E, 0x5A, 0x43)
GOLD = RGBColor(0xC4, 0x9A, 0x2B)
DARK = RGBColor(0x1A, 0x1A, 0x1A)
GREY = RGBColor(0x5A, 0x5A, 0x5A)
BG = RGBColor(0xF7, 0xF4, 0xEB)

SLIDES = [
    {"t": "ميزان  Mizan — الشارةُ محسوبة لا مُدّعاة",
     "s": "فحصٌ آلي لكل اقتباس من القرآن والحديث قبل أن يظهر أمام أي إنسان",
     "b": ["مسابقة الذكاء الاصطناعي في خدمة المحتوى الإسلامي 2026 · islamicaich.org",
           "المسار 04: أدوات المعرفة والتحقق — الثقة والسلامة العلمية شرطان لا يُستثنيان",
           "شعارنا: نُعلّم الآلة لتخدم الرسالة — والتزامنا دليلٌ لا ادعاء",
           "جرّبه بنفسك: bun run demo — بلا إنترنت ولا مفتاح"],
     "img": False},
    {"t": "1 · المشكلة: نموذجٌ يخترع نصًّا شرعيًّا بثقة تامة",
     "s": "هنا الوهم أخطر من الكذب: الخطأ يُصنع باسم دين، لا بخطأ في تجربة مستخدم",
     "b": ["نماذج أجنبية تهيمن على العربية، ومصادر إسلامية موثوقة قليلة جدًا",
           "الاقتباس المخترع لا يُفرّقه العام عن الحقيقي — والخطأ هنا يمسّ العبادة والاعتقاد",
           "في قياس 2026: دقة النماذج في التحقق من الحديث ≈ 8٪ — فلا تجعل النموذج حاكمًا",
           "أدوات التحقق بالتشابه تُعطي الحديث المخترع درجة عالية — وهذا خطر مُثبت بالتجربة (CWE-345)",
           "البحث العادي يقول «وجدناه» ولا يقول «هذه العبارة موجودة في هذا الكتاب تحديدًا»"],
     "img": False},
    {"t": "2 · الحل: أجِبْ أولًا، ثم أثبِتْ كل اقتباس على حدة",
     "s": "ثلاث شارات، كل واحدة محسوبة من النص نفسه — لا يطلبها النموذج ولا يوقّعها بشر",
     "b": ["مُتحقَّق: نص الاقتباس موجود حرفيًا داخل المصدر المذكور (بعد توحيد الكتابة) — ولا شيء غير ذلك",
           "مرفوض: المصدر موجود لكنه لا يحتوي الاقتباس — ويُعرض النصان معًا مع رابط المصدر",
           "غير قابل للتحقق: لا مصدر / نص فارغ / معرّف مجهول / انتهاء وقت — صفر دليل يعني صفر اعتماد",
           "الشارة تجيب عن سؤال واحد: هل هذا الاقتباس موجود في هذا المرجع؟ لا تدّعي صحة الحديث ولا درجته ولا فتوى"],
     "img": True},
    {"t": "3 · آلية العمل: ست خطوات وقاعدة توحيد كتابة واحدة",
     "s": "كل خطوة ترفض مبكرًا، والنتيجة الجيدة تأتي في آخر خطوة",
     "b": ["١ اقتباس؟  ٢ مصدر؟  ٣ ≤ ٣ مصادر؟  ٤ هل المعرّف يطابق سجلًّا؟  ٥ مطابقة حرفية صارمة؟  ٦ دليل مرفق؟",
           "قاعدة توحيد الكتابة في مكان واحد: الألف والهمزات والتاء المربوطة والتشكيل والتطويل",
           "وتُوحَّد الأرقام والترقيم والمسافات — بلا إضافة ولا حذف، فلا تُقارن كتابة مخترعة بالأصل",
           "textMatch مفتاح للمطابقة فقط ولا يُعرض أبدًا · والنص الذي يراه المستخدم هو textDisplay كما هو",
           "المُتحقَّق يعتمد على حزمة واحدة فقط (core) — بلا إنترنت ولا وقت حيّ ولا عشوائية ولا حكم من نموذج"],
     "img": False},
    {"t": "4 · الابتكار: لا تشابه على الشارة — حرفٌ واحد فقط",
     "s": "ليس اختيارًا تقنيًا فحسب — بل قياسٌ علمي وجد الثغرة",
     "b": ["لا تشابه ولا مسافة تحرير ولا نسبة ولا حدّ مسموح: الطريقة الوحيدة للوصول إلى «مُتحقَّق» هي المطابقة الحرفية",
           "النسبة التقريبية مجرد تشخيص جانبي يُعرض للاطلاع فقط، ولا تدخل في الشارة أبدًا",
           "الإثبات آلي لا وعود: G-1 يمنع أي مكتبة ممنوعة داخل المُتحقَّق · G-6 يفرض مكانًا واحدًا لبناءه — وأي طريق ثانٍ يُفشل الـCI",
           "المُتحقَّق لا يستورد كود العرض ولا أداة التشابه — الفصل مضمون ببوابة آلية لا بمراجعة بشرية",
           "الفارق عن RAG التقليدي: الاسترجاع يُجيب، والتحقّق يمنع أن يصير الوهم موثوقًا"],
     "img": False},
    {"t": "5 · الأدلة: تجربةٌ تُعيد الحساب، فإن لم تنجح تُلغى",
     "s": "أرقامٌ محسوبة بتشغيل فعلي، لا مكتوبة بيد أحد — والنتيجة السيئة عندنا صفر لا نسبة",
     "b": ["٢٠٠ حالة صحيحة مُتوقَّعة: نصٌّ مشكّل وبلا تشكيل ومُطوّل وأرقام — كلها تتحقق، والتوقّع مكتوب يدويًا",
           "٤٠ حالة تُقصد بالخطأ: كلمة تغيّرت، حرف انقلب، رقم تبدّل ← لا «مُتحقَّق» إطلاقًا. المطلوب صفر خطأ",
           "المولّد ممنوع من الوصول إلى المُتحقَّق — والاختبار يثبت هذا المنع، فالتوقّع يأتي من تشغيل فعلي لا من ورقة",
           "ومقارنة ببحث نصي مباشر (FTS5) على نفس الـ٤٠ حالة: كشف ٦٥٪ مقابل ١٠٠٪، وصفر اعتماد كاذب",
           "الاختبارات الآلية خضراء في كل حزمة + بوابات G-1…G-7 · وdemo يعمل في ثانية بلا مفتاح"],
     "img": False},
    {"t": "6 · الموثوقية والنسب: ما لا نملكه لا ندّعيه",
     "s": "باختصار: المصدر، وما يثبته الدليل، ومتى نرفض، ومتى نُحال إلى المختص",
     "b": ["27,234 سجلًا جاهزًا: القرآن (تنزيل عثماني) + 4 سنن + الموطأ — النسائي، أبو داود، ابن ماجه، الترمذي",
           "الدرجة كما يدعيها المصدر نفسه مع سببها، وإلا كانت null — لا نرفعها ولا نخمّنها",
           "١٥٬٠٢٦ صفًّا بلا درجة أُخفيت (٤١٫٧٪ من مصدر الحديث) — قبِلنا هذا النقص بدل تقديم ما لا يستند إلى دليل",
           "البخاري ومسلم مسجّلان «غائبَين» ولا يُعوَّضان — الاعتراف بالنقص جزء من الموثوقية",
           "لكل سجل ختمٌ رقمي متسلسل؛ تعارض الختم = خطأ واضح بلا شارة · والأسئلة تُحفظ بصيغة hash فقط"],
     "img": False},
]

FOOTER = "ميزان (mizan)  ·  الشارة محسوبة لا مُدّعاة  ·  islamicaich.org — المسار 04"


def pick_fonts():
    latin = os.path.join(FONTS, "arial.ttf")
    latin_b = os.path.join(FONTS, "arialbd.ttf")
    mono = os.path.join(FONTS, "cour.ttf")
    for p in (latin, latin_b, mono):
        if not os.path.exists(p):
            raise SystemExit("missing font: " + p)
    return latin, latin_b, mono


LATIN, LATIN_B, MONO = pick_fonts()
_reshaper = arabic_reshaper.ArabicReshaper({"delete_harakat": False, "support_ligatures": True})
# PIL cannot stack combining marks (no raqm), so the terminal screenshot
# renders unvocalized Arabic; the PDF keeps the harakat.
_reshaper_flat = arabic_reshaper.ArabicReshaper({"delete_harakat": True, "support_ligatures": True})


def disp(text):
    return get_display(_reshaper.reshape(text))


def disp_flat(text):
    return get_display(_reshaper_flat.reshape(text))


# ---------------------------------------------------------------- terminal png
def build_terminal_png():
    keep_sub = ("  1/2  ikhlas", "  2/2  fabricated-hadith")
    drop_pref = ("    relevance:", "correction:", "  shows:", "sources", "model ",
                 "transcript", "snapshot")
    lines = []
    with open(TERMINAL_SRC, encoding="utf-8", errors="replace") as fh:
        raw = [ln.rstrip("\n") for ln in fh]
    take = False
    for ln in raw:
        if ln.strip().startswith("mizan \u2014 offline demo") or ln.strip().startswith("no API key"):
            lines.append(ln.strip())
            continue
        if any(ln.startswith(k) for k in keep_sub):
            take = True
        elif ln.startswith("  \u2713"):
            take = False
            lines.append(ln)
            continue
        if not take:
            continue
        if any(ln.startswith(d) for d in drop_pref):
            continue
        stripped = ln.strip()
        if len(stripped) > 5 and set(stripped) <= set("\u2500\u2501\u2550- \u250c\u2510\u2514\u2518"):
            continue
        lines.append(ln)
    out = []
    for ln in lines:
        ln = ln.replace("\u2713", "\u00bb")
        while len(ln) > 84:
            out.append(ln[:84])
            ln = "        " + ln[84:]
        out.append(ln)

    font = ImageFont.truetype(MONO, 18)
    pad = 22
    lh = 24
    probe_w = 84 * 11
    img = Image.new("RGB", (probe_w + 2 * pad, lh * len(out) + 2 * pad), "#0d1117")
    dr = ImageDraw.Draw(img)
    y = pad
    for ln in out:
        shown = ln
        color = "#e6edf3"
        if "[VERIFIED]" in ln:
            color = "#3fb950"
        elif "[REJECTED]" in ln:
            color = "#f85149"
        elif "reached" in ln:
            color = "#3fb950"
        elif ln.startswith("mizan"):
            color = "#58a6ff"
        if any(ln.lstrip().startswith(p) for p in ("quoted:", "source:")):
            color = "#d2a8ff"
        if any("\u0600" <= c <= "\u06ff" for c in ln):
            shown = disp_flat(ln)
        dr.text((pad, y), shown, font=font, fill=color)
        y += lh
    bbox = img.getbbox()
    img = img.crop(bbox) if bbox else img
    img.save(TERMINAL_PNG)
    return TERMINAL_PNG


# ---------------------------------------------------------------- pptx
def _rtl_para(p, size, bold=False, color=DARK, align_right=True):
    if align_right:
        p.alignment = PP_ALIGN.RIGHT
    pPr = p._p.get_or_add_pPr()
    pPr.set("rtl", "1" if align_right else "0")
    for run in p.runs:
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = "Arial"
        rPr = run._r.get_or_add_rPr()
        for tag in ("a:ea", "a:cs"):
            el = rPr.makeelement(qn(tag), {"typeface": "Arial"})
            rPr.append(el)


def build_pptx(png):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    total = len(SLIDES)
    for idx, sl in enumerate(SLIDES):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        fill = s.background.fill
        fill.solid()
        fill.fore_color.rgb = BG
        bar = s.shapes.add_shape(1, 0, 0, Inches(13.333), Inches(0.09))
        bar.fill.solid()
        bar.fill.fore_color.rgb = GREEN
        bar.line.fill.background()

        num = s.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(0.7), Inches(0.4)).text_frame
        num.text = f"{idx + 1}/{total}"
        for p in num.paragraphs:
            for r in p.runs:
                r.font.size = Pt(11)
                r.font.color.rgb = GREY
                r.font.name = "Arial"

        text_w = Inches(7.4) if sl["img"] else Inches(11.9)
        text_left = Inches(5.4) if sl["img"] else Inches(0.7)

        tb = s.shapes.add_textbox(text_left, Inches(0.35), text_w, Inches(0.9)).text_frame
        tb.word_wrap = True
        p = tb.paragraphs[0]
        p.text = sl["t"]
        _rtl_para(p, 28, bold=True, color=GREEN)

        sb = s.shapes.add_textbox(text_left, Inches(1.25), text_w, Inches(0.7)).text_frame
        sb.word_wrap = True
        p = sb.paragraphs[0]
        p.text = sl["s"]
        _rtl_para(p, 16, color=DARK)

        rule = s.shapes.add_shape(1, text_left + text_w - Inches(1.1), Inches(2.0), Inches(1.1), Pt(4))
        rule.fill.solid()
        rule.fill.fore_color.rgb = GOLD
        rule.line.fill.background()

        bb = s.shapes.add_textbox(text_left, Inches(2.2), text_w, Inches(4.5)).text_frame
        bb.word_wrap = True
        for i, b in enumerate(sl["b"]):
            p = bb.paragraphs[0] if i == 0 else bb.add_paragraph()
            p.text = "•  " + b
            p.space_after = Pt(15)
            _rtl_para(p, 17 if sl["img"] else 19, color=DARK)

        if sl["img"]:
            iw, ih = Image.open(png).size
            aspect = iw / ih
            height_in = 4.4
            width_in = height_in * aspect
            if width_in > 6.0:
                width_in = 6.0
                height_in = width_in / aspect
            s.shapes.add_picture(png, Inches(0.5), Inches(2.2), width=Inches(width_in),
                                 height=Inches(height_in))

        ft = s.shapes.add_textbox(Inches(0.7), Inches(6.95), Inches(11.9), Inches(0.4)).text_frame
        p = ft.paragraphs[0]
        p.text = FOOTER
        p.alignment = PP_ALIGN.LEFT
        for r in p.runs:
            r.font.size = Pt(10)
            r.font.color.rgb = GREY
            r.font.name = "Arial"

    out = os.path.join(HERE, "mizan-idea-7slides-ar.pptx")
    return _save(prs, out)


# ---------------------------------------------------------------- pdf
def build_pdf(png):
    pdfmetrics.registerFont(TTFont("ArialAr", LATIN))
    pdfmetrics.registerFont(TTFont("ArialAr-Bold", LATIN_B))
    W, H = landscape(A4)
    final = os.path.join(HERE, "mizan-idea-7slides-ar.pdf")
    tmp = os.path.join(HERE, ".deck-ar.tmp.pdf")
    c = canvas.Canvas(tmp, pagesize=(W, H))
    total = len(SLIDES)

    def wrap(text, font, size, maxw):
        words = text.split(" ")
        lines, cur = [], ""
        for w in words:
            trial = (cur + " " + w).strip()
            if pdfmetrics.stringWidth(disp(trial), font, size) <= maxw or not cur:
                cur = trial
            else:
                lines.append(cur)
                cur = w
        if cur:
            lines.append(cur)
        return lines

    for idx, sl in enumerate(SLIDES):
        c.setFillColor(HexColor("#F7F4EB"))
        c.rect(0, 0, W, H, fill=1, stroke=0)
        c.setFillColor(HexColor("#0E5A43"))
        c.rect(0, H - 8, W, 8, fill=1, stroke=0)
        c.setFillColor(HexColor("#5A5A5A"))
        c.setFont("ArialAr", 9)
        c.drawString(34, H - 26, f"{idx + 1}/{total}")

        right = W - 40
        maxw = 400 if sl["img"] else 690

        c.setFillColor(HexColor("#0E5A43"))
        c.setFont("ArialAr-Bold", 22)
        y = H - 58
        for ln in wrap(sl["t"], "ArialAr-Bold", 22, maxw):
            c.drawRightString(right, y, disp(ln))
            y -= 27

        c.setFillColor(HexColor("#1A1A1A"))
        c.setFont("ArialAr", 13.5)
        for ln in wrap(sl["s"], "ArialAr", 13.5, maxw):
            c.drawRightString(right, y, disp(ln))
            y -= 18
        y -= 8

        c.setFillColor(HexColor("#C49A2B"))
        c.rect(right - 90, y + 6, 90, 4, fill=1, stroke=0)
        y -= 16

        c.setFont("ArialAr", 15)
        for b in sl["b"]:
            for j, ln in enumerate(wrap("\u2022  " + b, "ArialAr", 15, maxw)):
                c.setFillColor(HexColor("#1A1A1A"))
                prefix = "" if j == 0 else "    "
                c.drawRightString(right, y, disp(prefix + ln))
                y -= 19.5
            y -= 11

        if sl["img"]:
            iw, ih = Image.open(png).size
            target_w = 360
            target_h = target_w * ih / iw
            if target_h > 380:
                target_h = 380
                target_w = target_h * iw / ih
            img_y = 100 + (370 - target_h) / 2
            c.drawImage(png, 34, img_y, width=target_w, height=target_h)

        c.setFillColor(HexColor("#5A5A5A"))
        c.setFont("ArialAr", 9)
        c.drawString(34, 30, disp(FOOTER))
        c.showPage()
    c.save()
    with open(tmp, "rb") as fh:
        data = fh.read()
    os.remove(tmp)
    return _place(data, final)


def _save(prs, path):
    try:
        prs.save(path)
        return path
    except PermissionError:
        alt = path.replace(".pptx", "-v2.pptx")
        prs.save(alt)
        return alt


def _place(src_bytes, path):
    try:
        with open(path, "wb") as fh:
            fh.write(src_bytes)
        return path
    except PermissionError:
        alt = path.replace(".pdf", "-v2.pdf").replace(".pptx", "-v2.pptx")
        with open(alt, "wb") as fh:
            fh.write(src_bytes)
        return alt


def main():
    png = build_terminal_png()
    pptx = build_pptx(png)
    pdf = build_pdf(png)
    print("slides:", len(SLIDES))
    print("png:", png)
    print("pptx:", pptx)
    print("pdf:", pdf)


if __name__ == "__main__":
    main()
