"""Build an enriched simple-words mizan deck: modern Islamic UX, PPTX + PDF.

Design: deep emerald + gold on cream, badge pills, cards, star motif,
terminal mockup. Wording stays simple. Outputs stay far below 20 MB.
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from reportlab.lib.pagesizes import landscape, A4
from reportlab.lib.colors import HexColor
from reportlab.pdfgen import canvas

# Palette
EMERALD = RGBColor(0x0B, 0x3D, 0x2E)
EMERALD2 = RGBColor(0x0E, 0x5A, 0x43)
GOLD = RGBColor(0xC4, 0x9A, 0x2B)
GOLD_LT = RGBColor(0xE8, 0xD5, 0xA3)
CREAM = RGBColor(0xFA, 0xF6, 0xEC)
SAND = RGBColor(0xF1, 0xEA, 0xD8)
INK = RGBColor(0x1E, 0x2A, 0x26)
MUTED = RGBColor(0x5C, 0x6B, 0x66)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MINT_BG = RGBColor(0xDC, 0xFC, 0xE7)
MINT_TX = RGBColor(0x15, 0x80, 0x3D)
RED_BG = RGBColor(0xFE, 0xE2, 0xE2)
RED_TX = RGBColor(0xB9, 0x1C, 0x1C)
AMB_BG = RGBColor(0xFE, 0xF3, 0xC7)
AMB_TX = RGBColor(0xB4, 0x53, 0x09)
TERM = RGBColor(0x10, 0x18, 0x15)
TERM_LT = RGBColor(0x1E, 0x2E, 0x27)
HDR_SUB = RGBColor(0xD6, 0xE4, 0xDC)
CARD_LINE = RGBColor(0xE3, 0xD9, 0xC2)

W = Inches(13.333)
H = Inches(7.5)

try:
    from arabic_reshaper import reshape as _reshape
    from bidi.algorithm import get_display as _get_display
    from reportlab.pdfbase import pdfmetrics as _pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont as _TTFont
    _pdfmetrics.registerFont(_TTFont("Arial", "C:/Windows/Fonts/arial.ttf"))
    _pdfmetrics.registerFont(_TTFont("Arial-Bold", "C:/Windows/Fonts/arialbd.ttf"))
    HAS_ARABIC = True
except Exception:
    HAS_ARABIC = False

def _has_ar(s):
    return any("\u0600" <= ch <= "\u06FF" for ch in s)

def pdf_text(s):
    if not HAS_ARABIC or not _has_ar(s):
        return s
    return _get_display(_reshape(s))

def pdf_font(bold=False):
    if not HAS_ARABIC:
        return "Helvetica-Bold" if bold else "Helvetica"
    return "Arial-Bold" if bold else "Arial"


def _box(slide, shape, l, t, w, h, fill=None, line=None, lw=None):
    sp = slide.shapes.add_shape(shape, l, t, w, h)
    sp.line.fill.background()
    if fill is not None:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    else:
        sp.fill.background()
    if line is not None:
        sp.line.color.rgb = line
        sp.line.width = lw or Pt(1)
    sp.text_frame.word_wrap = True
    return sp


def _tb(slide, l, t, w, h):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tb.text_frame.word_wrap = True
    return tb.text_frame


def _para(tf, text, size, color, bold=False, align=PP_ALIGN.LEFT,
         first=False, space_after=Pt(4), font="Calibri"):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = font
    p.alignment = align
    p.space_after = space_after
    p.space_before = Pt(0)
    return p


def _stars(slide, cx, cy, n=4, step=0.34, size=0.13):
    for i in range(n):
        d = _box(slide, MSO_SHAPE.DIAMOND,
                 Inches(cx + i * step), Inches(cy), Inches(size), Inches(size),
                 fill=GOLD if i % 2 == 0 else GOLD_LT)
        d.rotation = 0.0


def _footer(slide, idx, total, dark=False):
    c = GOLD_LT if dark else MUTED
    tf = _tb(slide, Inches(0.6), Inches(7.02), Inches(10.5), Inches(0.35))
    _para(tf, "mizan (\u0645\u064a\u0632\u0627\u0646)  \u00b7  the badge is computed, not asserted",
          10, c, first=True, space_after=Pt(0))
    tf2 = _tb(slide, Inches(11.9), Inches(7.02), Inches(0.8), Inches(0.35))
    _para(tf2, f"{idx + 1}/{total}", 10, c, align=PP_ALIGN.RIGHT, first=True)


def _header(slide, kicker, title, subtitle, idx, total):
    _box(slide, MSO_SHAPE.RECTANGLE, 0, 0, W, Inches(1.62), fill=EMERALD)
    _box(slide, MSO_SHAPE.RECTANGLE, 0, Inches(1.62), W, Inches(0.055), fill=GOLD)
    _stars(slide, 11.55, 0.55, n=4)
    # kicker pill
    pill = _box(slide, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(0.22),
                Inches(2.35), Inches(0.34), fill=GOLD)
    tf = pill.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    _para(tf, kicker.upper(), 10, EMERALD, bold=True, align=PP_ALIGN.CENTER, first=True)
    tf2 = _tb(slide, Inches(0.6), Inches(0.62), Inches(10.4), Inches(0.55))
    _para(tf2, title, 27, WHITE, bold=True, first=True, space_after=Pt(0))
    tf3 = _tb(slide, Inches(0.6), Inches(1.10), Inches(10.4), Inches(0.45))
    _para(tf3, subtitle, 13.5, HDR_SUB, first=True)
    _footer(slide, idx, total)


def _card(slide, l, t, w, h, accent=None):
    c = _box(slide, MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h, fill=WHITE, line=CARD_LINE)
    if accent is not None:
        _box(slide, MSO_SHAPE.ROUNDED_RECTANGLE, l + Inches(0.22), t + Inches(0.18),
             Inches(0.55), Inches(0.07), fill=accent)
    return c


def _pill(slide, l, t, w, text, bg, fg, size=12):
    p = _box(slide, MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, Inches(0.42), fill=bg, line=fg, lw=Pt(1.25))
    tf = p.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    _para(tf, text, size, fg, bold=True, align=PP_ALIGN.CENTER, first=True)
    return p


def _chip(slide, l, t, w, h, text, size=12):
    c = _box(slide, MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h, fill=TERM, line=GOLD, lw=Pt(1))
    tf = c.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    _para(tf, text, size, GOLD_LT, font="Consolas", first=True, space_after=Pt(0))
    return c


prs = Presentation()
prs.slide_width = W
prs.slide_height = H
TOTAL = 9

# ---------- 0 · Cover ----------
s = prs.slides.add_slide(prs.slide_layouts[6])
_box(s, MSO_SHAPE.RECTANGLE, 0, 0, W, H, fill=EMERALD)
ring = _box(s, MSO_SHAPE.OVAL, Inches(9.2), Inches(-1.6), Inches(6.0), Inches(6.0),
            fill=None, line=GOLD)
ring.line.width = Pt(2.25)
ring2 = _box(s, MSO_SHAPE.OVAL, Inches(9.7), Inches(-1.1), Inches(5.0), Inches(5.0),
             fill=None, line=GOLD_LT)
ring2.line.width = Pt(1)
_stars(s, 0.6, 0.45, n=5)
tf = _tb(s, Inches(0.6), Inches(0.85), Inches(7.6), Inches(1.3))
_para(tf, "\u0645\u0650\u064a\u0632\u064e\u0627\u0646", 64, GOLD, bold=True, first=True, space_after=Pt(0))
tf2 = _tb(s, Inches(0.65), Inches(2.15), Inches(7.6), Inches(0.7))
_para(tf2, "mizan \u2014 the balance", 34, WHITE, bold=True, first=True, space_after=Pt(0))
tf3 = _tb(s, Inches(0.65), Inches(2.85), Inches(7.4), Inches(1.1))
_para(tf3, "An AI that answers Qur\u2019an & hadith questions \u2014 then proves every quote.", 16, HDR_SUB, first=True)
_para(tf3, "Computed badges. Shown evidence. No guessing.", 16, GOLD_LT, bold=True)
_pill(s, Inches(0.65), Inches(4.15), Inches(1.9), "\u2713 VERIFIED", MINT_BG, MINT_TX)
_pill(s, Inches(2.7), Inches(4.15), Inches(1.9), "\u2715 REJECTED", RED_BG, RED_TX)
_pill(s, Inches(4.75), Inches(4.15), Inches(2.2), "? UNVERIFIABLE", AMB_BG, AMB_TX)
_chip(s, Inches(0.65), Inches(4.85), Inches(5.4), Inches(0.55), "$  bun run demo    \u2192   VERIFIED + REJECTED in ~1s, offline")
tf4 = _tb(s, Inches(0.65), Inches(5.65), Inches(7.4), Inches(0.5))
_para(tf4, "Track 04 \u00b7 Knowledge & verification tools", 12.5, GOLD_LT, bold=True, first=True)
tfr = _tb(s, Inches(9.0), Inches(4.9), Inches(3.7), Inches(1.9))
_para(tfr, "27,234", 40, WHITE, bold=True, align=PP_ALIGN.RIGHT, first=True, space_after=Pt(0))
_para(tfr, "real records checked locally\nQur\u2019an + 4 Sunan + Muwatta", 13, HDR_SUB, align=PP_ALIGN.RIGHT)
_footer(s, 0, TOTAL, dark=True)

# ---------- 1 · Problem ----------
s = prs.slides.add_slide(prs.slide_layouts[6])
_box(s, MSO_SHAPE.RECTANGLE, 0, 0, W, H, fill=CREAM)
_header(s, "The problem", "AI makes up religious text \u2014 and it looks real",
        "Fluent words are not proof. A fake quote fools even careful readers.", 1, TOTAL)
cards = [
    ("AI writes fluently", "Models \"hallucinate\": they produce hadith that read perfectly \u2014 but do not exist.", "\u270e"),
    ("Fake looks real", "To a normal reader, an invented quote and a real one look identical.", "\u25c9"),
    ("Fuzzy checkers fail", "A fake-but-plausible hadith can score 97% on similarity. That is the hole (CWE-345).", "!"),
]
for i, (h, b, icon) in enumerate(cards):
    x = Inches(0.6 + i * 4.05)
    _card(s, x, Inches(2.25), Inches(3.75), Inches(2.9), accent=GOLD)
    tfi = _tb(s, x + Inches(0.22), Inches(2.55), Inches(0.5), Inches(0.5))
    _para(tfi, icon, 22, GOLD, bold=True, first=True)
    tfh = _tb(s, x + Inches(0.22), Inches(3.05), Inches(3.3), Inches(0.55))
    _para(tfh, h, 17, INK, bold=True, first=True)
    tfb = _tb(s, x + Inches(0.22), Inches(3.6), Inches(3.3), Inches(1.3))
    _para(tfb, b, 13.5, MUTED, first=True)
warn = _box(s, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(5.45), Inches(12.1), Inches(1.25),
            fill=RED_BG, line=RED_TX, lw=Pt(1.25))
tfw = warn.text_frame
tfw.vertical_anchor = MSO_ANCHOR.MIDDLE
_para(tfw, "Why it matters:  in worship & belief, a wrong quote is not a typo \u2014 it becomes false authority.",
      15, RED_TX, bold=True, align=PP_ALIGN.CENTER, first=True)

# ---------- 2 · Solution ----------
s = prs.slides.add_slide(prs.slide_layouts[6])
_box(s, MSO_SHAPE.RECTANGLE, 0, 0, W, H, fill=CREAM)
_header(s, "The solution", "Answer first \u2014 then prove every quote",
        "Three badges. One rule: only exact containment earns VERIFIED.", 2, TOTAL)
badges = [
    ("VERIFIED", MINT_BG, MINT_TX, "The quote is literally inside the source it cites.", EMERALD2),
    ("REJECTED", RED_BG, RED_TX, "The source exists \u2014 but the quote is NOT in it.", RED_TX),
    ("UNVERIFIABLE", AMB_BG, AMB_TX, "Cannot be checked: no citation, paraphrase, timeout\u2026", AMB_TX),
]
for i, (name, bgc, fgc, desc, acc) in enumerate(badges):
    x = Inches(0.6 + i * 4.05)
    _card(s, x, Inches(2.25), Inches(3.75), Inches(2.55), accent=acc)
    _pill(s, x + Inches(0.22), Inches(2.62), Inches(2.3), name, bgc, fgc, size=13)
    tfd = _tb(s, x + Inches(0.22), Inches(3.2), Inches(3.3), Inches(1.2))
    _para(tfd, desc, 14, INK, first=True)
ev = _box(s, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(5.1), Inches(12.1), Inches(1.6),
          fill=EMERALD, line=EMERALD)
tfe = _tb(s, Inches(1.0), Inches(5.28), Inches(11.3), Inches(1.3))
_para(tfe, "Every badge ships with its evidence  \u2014  quoted text  +  source row  +  source URL",
      15, GOLD_LT, bold=True, align=PP_ALIGN.CENTER, first=True, space_after=Pt(6))
_para(tfe, "If anything fails, the tool says so honestly. It never fakes certainty.", 13, WHITE,
      align=PP_ALIGN.CENTER)

# ---------- 3 · Workflow ----------
s = prs.slides.add_slide(prs.slide_layouts[6])
_box(s, MSO_SHAPE.RECTANGLE, 0, 0, W, H, fill=CREAM)
_header(s, "Workflow", "Six small checks \u2014 one path to VERIFIED",
        "Each claim walks the same steps. The first failure decides the badge.", 3, TOTAL)
steps = [
    ("1", "Quote?", "empty \u2192\nUNVERIFIABLE"),
    ("2", "Citation?", "none \u2192\nUNVERIFIABLE"),
    ("3", "\u22643 cites?", "cap the\nrest"),
    ("4", "Record\nfound?", "missing \u2192\nUNVERIFIABLE"),
    ("5", "Inside it?", "exact match \u2192\nVERIFIED"),
    ("6", "Proof?", "no proof \u2192\nno badge"),
]
sx, sw, gap = 0.6, 1.72, 0.32
for i, (n, h, b) in enumerate(steps):
    x = Inches(sx + i * (sw + gap))
    last = (i == 4)
    _card(s, x, Inches(2.25), Inches(sw), Inches(2.6), accent=GOLD if last else EMERALD2)
    num = _box(s, MSO_SHAPE.OVAL, x + Inches(0.55), Inches(2.5), Inches(0.62), Inches(0.62),
               fill=GOLD if last else EMERALD)
    tfn = num.text_frame
    tfn.vertical_anchor = MSO_ANCHOR.MIDDLE
    _para(tfn, n, 16, WHITE if not last else EMERALD, bold=True, align=PP_ALIGN.CENTER, first=True)
    tfh = _tb(s, x + Inches(0.1), Inches(3.22), Inches(sw - 0.2), Inches(0.9))
    _para(tfh, h, 14.5, INK, bold=True, align=PP_ALIGN.CENTER, first=True)
    tfb = _tb(s, x + Inches(0.1), Inches(3.95), Inches(sw - 0.2), Inches(0.8))
    _para(tfb, b, 11.5, MUTED, align=PP_ALIGN.CENTER, first=True)
    if i < 5:
        ta = _tb(s, x + Inches(sw) + Inches(0.02), Inches(3.25), Inches(gap - 0.04), Inches(0.5))
        _para(ta, "\u203a", 26, GOLD, bold=True, align=PP_ALIGN.CENTER, first=True)
note = _box(s, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(5.15), Inches(12.1), Inches(1.55),
            fill=SAND, line=GOLD, lw=Pt(1.25))
tfn2 = _tb(s, Inches(1.0), Inches(5.32), Inches(11.3), Inches(1.25))
_para(tfn2, "Normalization folds tashkeel, tatweel, alef forms & digits \u2014 but NEVER changes a letter.",
      14, INK, bold=True, align=PP_ALIGN.CENTER, first=True, space_after=Pt(4))
_para(tfn2, "So with or without \u0640\u0640\u0640\u062a\u0634\u0643\u064a\u0644 a true quote passes \u2014 and a one-word fake can never pass.",
      12.5, MUTED, align=PP_ALIGN.CENTER)

# ---------- 4 · Three ways ----------
s = prs.slides.add_slide(prs.slide_layouts[6])
_box(s, MSO_SHAPE.RECTANGLE, 0, 0, W, H, fill=CREAM)
_header(s, "Use it", "Three doors \u2014 same verifier behind each",
        "Command line, web page, or plug it into any AI assistant (MCP).", 4, TOTAL)
ways = [
    ("A \u00b7 Terminal", "Simplest. One second, offline.",
     ["$ bun run demo", "$ bun run ask \"\u2026\""], "No key \u00b7 no internet"),
    ("B \u00b7 Web page", "Click and see the badge.",
     ["$ bun run demo-server", "open localhost:3000"], "No code needed"),
    ("C \u00b7 MCP server", "For AI clients like Claude.",
     ["$ bun run mcp", "tool: verify (read-only)"], "Answers: verified /\nrejected / unverifiable"),
]
for i, (h, sub, cmds, foot) in enumerate(ways):
    x = Inches(0.6 + i * 4.05)
    _card(s, x, Inches(2.25), Inches(3.75), Inches(3.45), accent=EMERALD2)
    tfh = _tb(s, x + Inches(0.22), Inches(2.55), Inches(3.3), Inches(0.5))
    _para(tfh, h, 17, INK, bold=True, first=True, space_after=Pt(0))
    tfs = _tb(s, x + Inches(0.22), Inches(3.0), Inches(3.3), Inches(0.4))
    _para(tfs, sub, 12.5, MUTED, first=True)
    y = 3.5
    for cmd in cmds:
        _chip(s, x + Inches(0.22), Inches(y), Inches(3.3), Inches(0.5), cmd, size=11.5)
        y += 0.62
    tff = _tb(s, x + Inches(0.22), Inches(y + 0.08), Inches(3.3), Inches(0.8))
    _para(tff, foot, 12, EMERALD2, bold=True, first=True)

# ---------- 5 · Demo ----------
s = prs.slides.add_slide(prs.slide_layouts[6])
_box(s, MSO_SHAPE.RECTANGLE, 0, 0, W, H, fill=CREAM)
_header(s, "Live demo", "Real input \u2192 real output",
        "What a judge sees in the terminal in about one second.", 5, TOTAL)
term = _box(s, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(2.2), Inches(7.6), Inches(4.3),
            fill=TERM, line=EMERALD2, lw=Pt(1.5))
tft = _tb(s, Inches(0.95), Inches(2.45), Inches(6.9), Inches(3.9))
_para(tft, "$ bun run ask \"oneness of God?\"", 12, GOLD_LT, font="Consolas", first=True)
_para(tft, "[\u2713 VERIFIED]  ikhlas-1 \u2014 exact_containment", 12.5, WHITE, bold=True, font="Consolas")
_para(tft, "  quoted:  \u0642\u064f\u0644\u0652 \u0647\u064f\u0648\u064e \u0671\u0644\u0644\u064e\u0651\u0647\u064f \u0623\u064e\u062d\u064e\u062f\u064c", 13, MINT_BG, font="Consolas")
_para(tft, "  source:  quran 6222 \u2014 tanzil.net", 12, HDR_SUB, font="Consolas")
_para(tft, "[\u2715 REJECTED]  one-word-changed hadith", 12.5, RED_BG, bold=True, font="Consolas")
_para(tft, "  + real nearby hadith shown as proof", 12, HDR_SUB, font="Consolas")
_para(tft, "[? UNVERIFIABLE]  paraphrase \u2014 reason named", 12, AMB_BG, font="Consolas")
for i, (t, d) in enumerate([
    ("Same words?", "Containment compares folded text \u2014 quotes with or without tashkeel both pass."),
    ("Changed a word?", "Fold differs \u2192 REJECTED. No 97% escape hatch."),
    ("Only a meaning?", "Paraphrase \u2192 UNVERIFIABLE. Honest, not accused."),
]):
    x = Inches(8.5)
    y = Inches(2.2 + i * 1.5)
    _card(s, x, y, Inches(4.2), Inches(1.32), accent=GOLD)
    tfq = _tb(s, x + Inches(0.22), y + Inches(0.2), Inches(3.76), Inches(1.0))
    _para(tfq, t, 14.5, INK, bold=True, first=True, space_after=Pt(2))
    _para(tfq, d, 12, MUTED)

# ---------- 6 · Trust ----------
s = prs.slides.add_slide(prs.slide_layouts[6])
_box(s, MSO_SHAPE.RECTANGLE, 0, 0, W, H, fill=CREAM)
_header(s, "Trust", "Rules no one can bypass",
        "Not promises \u2014 machine-checked gates that fail the build.", 6, TOTAL)
trust = [
    ("Exact only", "No similarity, no percentages, no embedding path to VERIFIED."),
    ("One road", "A single VERIFIED construction site. CI fails if a second appears (G-6)."),
    ("Grades stay honest", "Stored exactly as the dataset says \u2014 or null. Never invented."),
    ("Private by design", "Logs carry hashes, never your question. Attested ledger."),
]
for i, (h, b) in enumerate(trust):
    col, row = i % 2, i // 2
    x = Inches(0.6 + col * 6.2)
    y = Inches(2.25 + row * 1.95)
    _card(s, x, y, Inches(5.9), Inches(1.7), accent=EMERALD2)
    tick = _box(s, MSO_SHAPE.OVAL, x + Inches(0.25), y + Inches(0.3), Inches(0.55), Inches(0.55),
                fill=MINT_BG, line=MINT_TX, lw=Pt(1.25))
    tftk = tick.text_frame
    tftk.vertical_anchor = MSO_ANCHOR.MIDDLE
    _para(tftk, "\u2713", 16, MINT_TX, bold=True, align=PP_ALIGN.CENTER, first=True)
    tfh = _tb(s, x + Inches(1.0), y + Inches(0.22), Inches(4.6), Inches(0.5))
    _para(tfh, h, 16, INK, bold=True, first=True)
    tfb = _tb(s, x + Inches(1.0), y + Inches(0.68), Inches(4.6), Inches(0.85))
    _para(tfb, b, 12.5, MUTED, first=True)

# ---------- 7 · Limits ----------
s = prs.slides.add_slide(prs.slide_layouts[6])
_box(s, MSO_SHAPE.RECTANGLE, 0, 0, W, H, fill=CREAM)
_header(s, "Honesty", "What we do NOT claim",
        "An over-claim here would be a religious error, not a marketing one.", 7, TOTAL)
lims = [
    "No Bukhari / Muslim yet \u2014 the biggest gap, stated upfront.",
    "Paraphrase \u21d2 UNVERIFIABLE, never VERIFIED (26 cases ruled by humans).",
    "Default mode replays a labelled recording; live mode needs your API key.",
    "It proves the QUOTE exists \u2014 not that the interpretation is correct.",
    "A hadith grade is the dataset\u2019s grade, never our ruling.",
]
box = _box(s, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(2.2), Inches(12.1), Inches(4.25),
           fill=AMB_BG, line=GOLD, lw=Pt(1.5))
for i, t in enumerate(lims):
    tfi = _tb(s, Inches(1.1), Inches(2.5 + i * 0.78), Inches(11.2), Inches(0.7))
    _para(tfi, f"\u25c7  {t}", 14, INK, first=True)

# ---------- 8 · Closing ----------
s = prs.slides.add_slide(prs.slide_layouts[6])
_box(s, MSO_SHAPE.RECTANGLE, 0, 0, W, H, fill=EMERALD)
_box(s, MSO_SHAPE.RECTANGLE, 0, Inches(7.0), W, Inches(0.5), fill=GOLD)
_stars(s, 11.55, 0.5, n=4)
tf = _tb(s, Inches(0.6), Inches(0.6), Inches(12.1), Inches(0.5))
_para(tf, "REMEMBER THREE LINES", 12, GOLD_LT, bold=True, align=PP_ALIGN.CENTER, first=True)
summ = [
    ("Problem", "AI invents hadith that look real."),
    ("Solution", "mizan checks each quote against the real source \u2014 exact match only."),
    ("Workflow", "ask \u2192 AI answers \u2192 every claim gets a computed badge + evidence."),
]
for i, (h, b) in enumerate(summ):
    y = Inches(1.35 + i * 1.28)
    _pill(s, Inches(1.1), y, Inches(2.2), h.upper(), GOLD, EMERALD, size=13)
    tfb = _tb(s, Inches(3.6), y - Inches(0.02), Inches(8.6), Inches(0.6))
    _para(tfb, b, 17, WHITE, first=True)
_chip(s, Inches(3.9), Inches(5.45), Inches(5.5), Inches(0.6),
      "$ bun install  &&  bun run demo", size=13)
tfa = _tb(s, Inches(0.6), Inches(6.15), Inches(12.1), Inches(0.6))
_para(tfa, "\u0646\u064f\u0639\u0644\u0650\u0651\u0645 \u0627\u0644\u0622\u0644\u0629 .. \u0644\u0650\u062a\u062e\u062f\u0645 \u0627\u0644\u0631\u0633\u0627\u0644\u0629 \u2014 with proof.",
      14, GOLD_LT, align=PP_ALIGN.CENTER, first=True)
_footer(s, 8, TOTAL, dark=True)

prs.save("submission/mizan-simple.pptx")

# ---------------- PDF mirror (enriched) ----------------
EM_H = HexColor("#0B3D2E")
GD_H = HexColor("#C49A2B")
CR_H = HexColor("#FAF6EC")
PDF_SLIDES = [
 ("START", "mizan (\u0645\u064a\u0632\u0627\u0646) \u2014 the balance",
  "An AI that answers Qur\u2019an & hadith questions \u2014 then proves every quote.",
  ["\u2713 VERIFIED \u00b7 \u2715 REJECTED \u00b7 ? UNVERIFIABLE \u2014 computed, never guessed",
   "$ bun run demo \u2192 VERIFIED + REJECTED in ~1s, offline, no key",
   "27,234 real records \u00b7 Qur\u2019an + 4 Sunan + Muwatta \u00b7 Track 04"]),
 ("THE PROBLEM", "AI makes up religious text \u2014 and it looks real",
  "Fluent words are not proof.",
  ["AI hallucinates fluent hadith that does not exist",
   "Fake and real quotes look identical to readers",
   "Fuzzy checkers fail: fake-but-plausible hadith \u2248 97% (CWE-345)",
   "In worship & belief, a wrong quote is not a typo"]),
 ("THE SOLUTION", "Answer first \u2014 then prove every quote",
  "Three badges. One rule: exact containment earns VERIFIED.",
  ["VERIFIED \u2192 quote literally inside the cited source",
   "REJECTED \u2192 source exists, quote NOT in it (+ real nearby text shown)",
   "UNVERIFIABLE \u2192 no citation / paraphrase / timeout (reason named)",
   "Every badge ships with: quoted text + source row + source URL"]),
 ("WORKFLOW", "Six small checks \u2014 one path to VERIFIED",
  "First failure decides the badge.",
  ["1 Quote? \u00b7 2 Citation? \u00b7 3 \u22643 cites? \u00b7 4 Record found?",
   "5 Inside it? exact folded match \u2192 VERIFIED \u00b7 6 Proof present?",
   "Folds tashkeel/tatweel/alef/digits \u2014 NEVER changes a letter",
   "True quote (with or without tashkeel) passes \u00b7 one-word fake never does"]),
 ("USE IT", "Three doors \u2014 same verifier",
  "CLI, web page, MCP server.",
  ["A \u00b7 Terminal: bun run demo (offline, ~1s) \u00b7 bun run ask \"\u2026\"",
   "B \u00b7 Web: bun run demo-server \u2192 localhost:3000",
   "C \u00b7 MCP: bun run mcp \u2192 read-only tool: verify",
   "Output per claim: verified / rejected / unverifiable"]),
 ("LIVE DEMO", "Real input \u2192 real output",
  "About one second in the terminal.",
  ["IN: bun run ask \"oneness of God?\"",
   "OUT: [\u2713 VERIFIED] ikhlas-1 \u00b7 quoted: \u0642\u064f\u0644\u0652 \u0647\u064f\u0648\u064e \u0627\u0644\u0644\u0647 \u0623\u062d\u062f \u00b7 quran 6222",
   "OUT: [\u2715 REJECTED] one-word-changed hadith + real nearby text",
   "OUT: [? UNVERIFIABLE] paraphrase \u2014 reason named"]),
 ("TRUST", "Rules no one can bypass",
  "Machine-checked gates, not promises.",
  ["Exact containment only \u2014 no similarity, no percentages",
   "Single VERIFIED site; CI fails if a second appears (G-6)",
   "Grades stored as dataset says, or null \u2014 never invented",
   "Logs carry hashes, never your question"]),
 ("HONESTY", "What we do NOT claim",
  "An over-claim here is a religious error.",
  ["No Bukhari / Muslim yet \u2014 biggest gap, stated upfront",
   "Paraphrase \u21d2 UNVERIFIABLE (26 human-ruled cases)",
   "Default = labelled replay; live needs your API key",
   "Proves the QUOTE exists \u2014 not the interpretation",
   "Grade = dataset\u2019s grade, never our ruling"]),
 ("REMEMBER", "Problem \u00b7 Solution \u00b7 Workflow",
  "\u0646\u0639\u0644\u0645 \u0627\u0644\u0622\u0644\u0629 .. \u0644\u062a\u062e\u062f\u0645 \u0627\u0644\u0631\u0633\u0627\u0644\u0629 \u2014 with proof.",
  ["Problem: AI invents hadith that look real",
   "Solution: check each quote \u2014 exact match only",
   "Workflow: ask \u2192 answer \u2192 computed badge + evidence",
   "Try: bun install && bun run demo"]),
]

c = canvas.Canvas("submission/mizan-simple.pdf", pagesize=landscape(A4))
PW, PH = landscape(A4)
for i, (kick, t, sub, bullets) in enumerate(PDF_SLIDES):
    c.setFillColor(CR_H); c.rect(0, 0, PW, PH, fill=1, stroke=0)
    c.setFillColor(EM_H); c.rect(0, PH - 120, PW, 120, fill=1, stroke=0)
    c.setFillColor(GD_H); c.rect(0, PH - 124, PW, 4, fill=1, stroke=0)
    c.setFillColor(GD_H); c.setFont(pdf_font(True), 9)
    c.drawString(36, PH - 40, pdf_text(kick))
    c.setFillColor(HexColor("#FFFFFF")); c.setFont(pdf_font(True), 20)
    c.drawString(36, PH - 68, pdf_text(t))
    c.setFillColor(HexColor("#D6E4DC")); c.setFont(pdf_font(), 11)
    c.drawString(36, PH - 90, pdf_text(sub))
    y = PH - 165
    for b in bullets:
        c.setFillColor(HexColor("#FFFFFF"))
        c.roundRect(36, y - 8, PW - 72, 30, 5, fill=1, stroke=0)
        c.setFillColor(GD_H); c.rect(36, y - 8, 5, 30, fill=1, stroke=0)
        c.setFillColor(HexColor("#1E2A26")); c.setFont(pdf_font(), 11)
        c.drawString(52, y + 2, pdf_text(b))
        y -= 40
    c.setFont(pdf_font(), 8); c.setFillColor(HexColor("#5C6B66"))
    c.drawString(36, 26, pdf_text("mizan (\u0645\u064a\u0632\u0627\u0646) \u00b7 the badge is computed, not asserted"))
    c.drawRightString(PW - 36, 26, f"{i + 1}/{len(PDF_SLIDES)}")
    c.showPage()
c.save()
print("done")
